from flask import Flask, jsonify, request
from flask_cors import CORS
import sqlite3
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_CEILING
import csv
import io
import os
from pathlib import Path

from werkzeug.security import check_password_hash, generate_password_hash

app = Flask(__name__)
CORS(app)

BASE_DIR = Path(__file__).resolve().parent
DB_NAME = Path(
    os.environ.get(
        "ALMADA2_DB_PATH",
        BASE_DIR / "data" / "almada2.db"
    )
)

DB_NAME.parent.mkdir(parents=True, exist_ok=True)

ADMIN_INICIAL = os.environ.get("ALMADA2_ADMIN_USER", "admin")
CLAVE_ADMIN_INICIAL = os.environ.get(
    "ALMADA2_ADMIN_PASSWORD",
    "almada2-dev"
)


def conectar_db():
    conexion = sqlite3.connect(
        str(DB_NAME),
        timeout=30,
        check_same_thread=False
    )
    conexion.row_factory = sqlite3.Row
    conexion.execute("PRAGMA foreign_keys = ON")
    return conexion


def fila_a_diccionario(fila):
    if fila is None:
        return None
    return dict(fila)


def limpiar_documento(valor):
    return "".join(caracter for caracter in str(valor or "") if caracter.isdigit())


def obtener_dni_desde_documento(valor):
    documento = limpiar_documento(valor)

    # CUIL/CUIT argentino: 11 dígitos. Ejemplo: 20 30123456 7
    if len(documento) == 11:
        return documento[2:10]

    # DNI común: últimos 8 dígitos
    if len(documento) >= 8:
        return documento[-8:]

    return documento


def calcular_precios(precio_base, bonif1, bonif2, bonif3, iva, flete, ganancia):
    precio = float(precio_base or 0)

    precio = precio * (1 - float(bonif1 or 0) / 100)
    precio = precio * (1 - float(bonif2 or 0) / 100)
    precio = precio * (1 - float(bonif3 or 0) / 100)

    precio = precio * (1 + float(iva or 0) / 100)
    precio = precio * (1 + float(flete or 0) / 100)

    precio_final_compra = round(precio, 2)
    precio_venta = round(precio_final_compra * (1 + float(ganancia or 0) / 100), 2)

    return precio_final_compra, precio_venta


def redondear_aumento_hacia_arriba_10(precio_actual, porcentaje):
    """Aplica el porcentaje y redondea hacia arriba al múltiplo de $10."""
    precio = Decimal(str(precio_actual or 0))
    aumento = Decimal(str(porcentaje))
    precio_calculado = precio * (Decimal("1") + aumento / Decimal("100"))
    precio_redondeado = (
        precio_calculado / Decimal("10")
    ).to_integral_value(rounding=ROUND_CEILING) * Decimal("10")
    return float(precio_redondeado)

def limpiar_texto(valor):
    return str(valor or "").strip()


def limpiar_numero(valor, defecto=0):
    if valor is None:
        return defecto

    texto = str(valor).strip()

    if texto == "":
        return defecto

    texto = texto.replace("$", "").replace(" ", "")

    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", ".")

    try:
        return float(texto)
    except ValueError:
        return defecto


def normalizar_columna(nombre):
    texto = str(nombre or "").strip().lower()
    reemplazos = {
        "á": "a",
        "é": "e",
        "í": "i",
        "ó": "o",
        "ú": "u",
        "ñ": "n"
    }

    for original, reemplazo in reemplazos.items():
        texto = texto.replace(original, reemplazo)

    texto = texto.replace(" ", "_").replace("-", "_")
    return texto


def obtener_valor(fila, *nombres):
    for nombre in nombres:
        clave = normalizar_columna(nombre)

        if clave in fila:
            return fila.get(clave)

    return ""


def leer_archivo_importacion(archivo):
    nombre = archivo.filename.lower()

    if nombre.endswith(".csv"):
        contenido = archivo.read().decode("utf-8-sig")
        muestra = contenido[:2048]

        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=";,")
            separador = dialecto.delimiter
        except Exception:
            separador = ";"

        lector = csv.DictReader(io.StringIO(contenido), delimiter=separador)

        filas = []
        for fila in lector:
            fila_normalizada = {}

            for clave, valor in fila.items():
                fila_normalizada[normalizar_columna(clave)] = valor

            filas.append(fila_normalizada)

        return filas

    if nombre.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(archivo, data_only=True)
        hoja = workbook.active

        encabezados = []
        filas = []

        for celda in hoja[1]:
            encabezados.append(normalizar_columna(celda.value))

        for row in hoja.iter_rows(min_row=2, values_only=True):
            fila = {}

            for index, valor in enumerate(row):
                if index < len(encabezados):
                    fila[encabezados[index]] = valor

            filas.append(fila)

        return filas

    raise ValueError("Formato no permitido. Usá CSV o XLSX.")

def generar_numero_cliente(cursor):
    cursor.execute("SELECT numero_cliente FROM clientes")
    clientes = cursor.fetchall()

    numero_mayor = 0

    for cliente in clientes:
        numero = cliente["numero_cliente"]

        if numero and str(numero).startswith("C"):
            try:
                numero_limpio = int(str(numero).replace("C", ""))
                if numero_limpio > numero_mayor:
                    numero_mayor = numero_limpio
            except ValueError:
                pass

    nuevo_numero = numero_mayor + 1
    return f"C{nuevo_numero:04d}"


def crear_tablas():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            usuario TEXT NOT NULL UNIQUE,
            contrasena TEXT NOT NULL,
            rol TEXT NOT NULL,
            activo INTEGER NOT NULL DEFAULT 1
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS clientes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            numero_cliente TEXT NOT NULL UNIQUE,
            nombre TEXT NOT NULL,
            direccion TEXT,
            localidad TEXT,
            telefono TEXT,
            email TEXT,
            observaciones TEXT,
            saldo_actual REAL DEFAULT 0
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS proveedores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            telefono TEXT,
            direccion TEXT,
            localidad TEXT,
            contacto TEXT,
            observaciones TEXT
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS productos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            proveedor_id INTEGER,
            codigo TEXT,
            nombre TEXT NOT NULL,
            descripcion TEXT,
            categoria TEXT,
            foto TEXT,
            cantidad_caja INTEGER DEFAULT 1,
            precio_unidad REAL DEFAULT 0,
            precio_caja REAL DEFAULT 0,
            precio_venta REAL DEFAULT 0,
            stock INTEGER DEFAULT 0,
            FOREIGN KEY (proveedor_id) REFERENCES proveedores(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS producto_variantes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            producto_id INTEGER NOT NULL,
            nombre_variante TEXT NOT NULL DEFAULT '',
            codigo TEXT DEFAULT '',
            cantidad_caja INTEGER DEFAULT 1,
            precio_unidad REAL DEFAULT 0,
            precio_caja REAL DEFAULT 0,
            precio_venta REAL DEFAULT 0,
            stock INTEGER DEFAULT 0,
            activo INTEGER DEFAULT 1,
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        )
    """)

    cursor.execute("""
        CREATE INDEX IF NOT EXISTS idx_variantes_producto
        ON producto_variantes(producto_id)
    """)

    cursor.execute("""
        CREATE INDEX IF NOT EXISTS idx_variantes_codigo
        ON producto_variantes(codigo)
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS precios_proveedor (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            proveedor_id INTEGER NOT NULL,
            producto_id INTEGER NOT NULL,
            precio_base REAL DEFAULT 0,
            bonificacion_1 REAL DEFAULT 0,
            bonificacion_2 REAL DEFAULT 0,
            bonificacion_3 REAL DEFAULT 0,
            iva REAL DEFAULT 21,
            flete REAL DEFAULT 0,
            precio_final_compra REAL DEFAULT 0,
            ganancia REAL DEFAULT 0,
            precio_venta REAL DEFAULT 0,
            fecha_actualizacion TEXT,
            FOREIGN KEY (proveedor_id) REFERENCES proveedores(id),
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS cuentas_corrientes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cliente_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            tipo TEXT NOT NULL,
            descripcion TEXT,
            comprobante TEXT,
            debe REAL DEFAULT 0,
            haber REAL DEFAULT 0,
            saldo REAL DEFAULT 0,
            medio_pago TEXT,
            FOREIGN KEY (cliente_id) REFERENCES clientes(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS pagos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cliente_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            monto REAL NOT NULL,
            medio_pago TEXT NOT NULL,
            comprobante TEXT,
            observaciones TEXT,
            FOREIGN KEY (cliente_id) REFERENCES clientes(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS saldos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cliente_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            descripcion TEXT,
            monto_inicial REAL NOT NULL,
            monto_pagado REAL DEFAULT 0,
            saldo_restante REAL NOT NULL,
            estado TEXT DEFAULT 'Pendiente',
            FOREIGN KEY (cliente_id) REFERENCES clientes(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS movimientos_stock (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            producto_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            tipo TEXT NOT NULL,
            cantidad INTEGER NOT NULL,
            descripcion TEXT,
            stock_resultante INTEGER,
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        )
    """)


    cursor.execute("""
        CREATE TABLE IF NOT EXISTS vendedores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            usuario_id INTEGER NOT NULL UNIQUE,
            nombre TEXT NOT NULL,
            telefono TEXT,
            zona TEXT,
            observaciones TEXT,
            activo INTEGER NOT NULL DEFAULT 1,
            fecha_alta TEXT NOT NULL,
            FOREIGN KEY (usuario_id) REFERENCES usuarios(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS stock_viaje (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            vendedor_id INTEGER NOT NULL,
            producto_id INTEGER NOT NULL,
            cantidad INTEGER NOT NULL DEFAULT 0,
            actualizado_en TEXT,
            UNIQUE (vendedor_id, producto_id),
            FOREIGN KEY (vendedor_id) REFERENCES vendedores(id),
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS ventas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            vendedor_id INTEGER NOT NULL,
            cliente_id INTEGER,
            fecha TEXT NOT NULL,
            forma_pago TEXT NOT NULL,
            total REAL NOT NULL DEFAULT 0,
            observaciones TEXT,
            FOREIGN KEY (vendedor_id) REFERENCES vendedores(id),
            FOREIGN KEY (cliente_id) REFERENCES clientes(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS venta_detalles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            venta_id INTEGER NOT NULL,
            producto_id INTEGER NOT NULL,
            cantidad INTEGER NOT NULL,
            precio_unitario REAL NOT NULL,
            subtotal REAL NOT NULL,
            FOREIGN KEY (venta_id) REFERENCES ventas(id),
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS movimientos_stock_viaje (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            vendedor_id INTEGER NOT NULL,
            producto_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            tipo TEXT NOT NULL,
            cantidad INTEGER NOT NULL,
            descripcion TEXT,
            stock_viaje_resultante INTEGER NOT NULL DEFAULT 0,
            stock_central_resultante INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY (vendedor_id) REFERENCES vendedores(id),
            FOREIGN KEY (producto_id) REFERENCES productos(id)
        )
    """)

    cursor.execute("""
        CREATE INDEX IF NOT EXISTS idx_stock_viaje_vendedor
        ON stock_viaje(vendedor_id)
    """)

    cursor.execute("""
        CREATE INDEX IF NOT EXISTS idx_ventas_vendedor
        ON ventas(vendedor_id)
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS aumentos_precios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            modo TEXT NOT NULL,
            porcentaje REAL NOT NULL,
            criterio TEXT NOT NULL,
            cantidad_productos INTEGER NOT NULL DEFAULT 0,
            cantidad_precios INTEGER NOT NULL DEFAULT 0,
            redondeo TEXT NOT NULL DEFAULT 'Hacia arriba a $10'
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS aumento_precio_detalles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            aumento_id INTEGER NOT NULL,
            producto_id INTEGER,
            producto TEXT NOT NULL,
            variante_id INTEGER,
            variante TEXT,
            codigo TEXT,
            proveedor TEXT,
            categoria TEXT,
            precio_anterior REAL NOT NULL,
            precio_nuevo REAL NOT NULL,
            FOREIGN KEY (aumento_id) REFERENCES aumentos_precios(id)
                ON DELETE CASCADE
        )
    """)

    cursor.execute("""
        CREATE INDEX IF NOT EXISTS idx_aumentos_precios_fecha
        ON aumentos_precios(fecha)
    """)

    cursor.execute("""
        CREATE INDEX IF NOT EXISTS idx_aumento_detalles_aumento
        ON aumento_precio_detalles(aumento_id)
    """)

    conexion.commit()
    conexion.close()


def agregar_columna_si_no_existe(tabla, columna, definicion):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute(f"PRAGMA table_info({tabla})")
    columnas = [fila["name"] for fila in cursor.fetchall()]

    if columna not in columnas:
        cursor.execute(f"ALTER TABLE {tabla} ADD COLUMN {columna} {definicion}")

    conexion.commit()
    conexion.close()


def asegurar_columnas():
    agregar_columna_si_no_existe("clientes", "saldo_actual", "REAL DEFAULT 0")

    agregar_columna_si_no_existe("productos", "foto_origen", "TEXT DEFAULT ''")
    agregar_columna_si_no_existe("productos", "foto_palabra_clave", "TEXT DEFAULT ''")


def cargar_datos_iniciales():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("SELECT id FROM usuarios WHERE usuario = ?", (ADMIN_INICIAL,))
    admin_existente = cursor.fetchone()

    if not admin_existente:
        cursor.execute("""
            INSERT INTO usuarios (usuario, contrasena, rol, activo)
            VALUES (?, ?, ?, 1)
        """, (
            ADMIN_INICIAL,
            generate_password_hash(CLAVE_ADMIN_INICIAL),
            "admin"
        ))

    conexion.commit()
    conexion.close()


@app.route("/")
def inicio():
    return jsonify({
        "mensaje": "API de Almada 2 funcionando correctamente"
    })


@app.route("/api/login", methods=["POST"])
def login():
    datos = request.json or {}

    usuario = datos.get("usuario", "").strip()
    contrasena = datos.get("contrasena", "").strip()

    if not usuario or not contrasena:
        return jsonify({
            "mensaje": "Usuario y contraseña son obligatorios"
        }), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT id, usuario, contrasena, rol, activo
        FROM usuarios
        WHERE usuario = ?
    """, (usuario,))

    usuario_encontrado = cursor.fetchone()

    if (
        not usuario_encontrado
        or not usuario_encontrado["activo"]
        or not check_password_hash(
            usuario_encontrado["contrasena"],
            contrasena
        )
    ):
        conexion.close()
        return jsonify({
            "mensaje": "Usuario o contraseña incorrectos"
        }), 401

    respuesta = {
        "mensaje": "Login correcto",
        "usuario_id": usuario_encontrado["id"],
        "usuario": usuario_encontrado["usuario"],
        "rol": usuario_encontrado["rol"]
    }

    if usuario_encontrado["rol"] == "vendedor":
        cursor.execute("""
            SELECT id, nombre, activo
            FROM vendedores
            WHERE usuario_id = ?
        """, (usuario_encontrado["id"],))
        vendedor = cursor.fetchone()

        if not vendedor or not vendedor["activo"]:
            conexion.close()
            return jsonify({
                "mensaje": "El vendedor no está habilitado"
            }), 403

        respuesta["vendedor_id"] = vendedor["id"]
        respuesta["nombre"] = vendedor["nombre"]

    conexion.close()
    return jsonify(respuesta)


@app.route("/api/resumen", methods=["GET"])
def resumen():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("SELECT COUNT(*) FROM productos")
    total_productos = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM productos WHERE stock <= 5")
    stock_bajo = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM proveedores")
    total_proveedores = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM clientes")
    total_clientes = cursor.fetchone()[0]

    cursor.execute("SELECT SUM(saldo_actual) FROM clientes")
    saldo_clientes = cursor.fetchone()[0] or 0

    cursor.execute("SELECT COUNT(*) FROM vendedores WHERE activo = 1")
    total_vendedores = cursor.fetchone()[0]

    cursor.execute("SELECT COALESCE(SUM(cantidad), 0) FROM stock_viaje")
    unidades_en_viaje = cursor.fetchone()[0] or 0

    mes_actual = datetime.now().strftime("%Y-%m")
    cursor.execute("""
        SELECT COALESCE(SUM(total), 0)
        FROM ventas
        WHERE substr(fecha, 1, 7) = ?
    """, (mes_actual,))
    ventas_mes = cursor.fetchone()[0] or 0

    conexion.close()

    return jsonify({
        "productos_cargados": total_productos,
        "stock_bajo": stock_bajo,
        "proveedores": total_proveedores,
        "clientes": total_clientes,
        "saldo_clientes": saldo_clientes,
        "vendedores": total_vendedores,
        "unidades_en_viaje": unidades_en_viaje,
        "ventas_mes": ventas_mes
    })


@app.route("/api/proveedores", methods=["GET"])
def listar_proveedores():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("SELECT * FROM proveedores ORDER BY nombre ASC")
    proveedores = cursor.fetchall()

    conexion.close()

    return jsonify([dict(proveedor) for proveedor in proveedores])


@app.route("/api/proveedores", methods=["POST"])
def agregar_proveedor():
    datos = request.json or {}

    nombre = datos.get("nombre", "").strip()

    if not nombre:
        return jsonify({"mensaje": "El nombre del proveedor es obligatorio"}), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        INSERT INTO proveedores (nombre, telefono, direccion, localidad, contacto, observaciones)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (
        nombre,
        datos.get("telefono", ""),
        datos.get("direccion", ""),
        datos.get("localidad", ""),
        datos.get("contacto", ""),
        datos.get("observaciones", "")
    ))

    conexion.commit()
    proveedor_id = cursor.lastrowid
    conexion.close()

    return jsonify({
        "mensaje": "Proveedor agregado correctamente",
        "id": proveedor_id
    }), 201


@app.route("/api/proveedores/<int:proveedor_id>/productos", methods=["GET"])
def productos_por_proveedor(proveedor_id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            pp.id AS precio_id,
            p.id AS producto_id,
            p.codigo,
            p.nombre AS producto,
            p.categoria,
            p.stock,
            pp.precio_base,
            pp.bonificacion_1,
            pp.bonificacion_2,
            pp.bonificacion_3,
            pp.iva,
            pp.flete,
            pp.precio_final_compra,
            pp.ganancia,
            pp.precio_venta,
            pp.fecha_actualizacion
        FROM precios_proveedor pp
        INNER JOIN productos p ON pp.producto_id = p.id
        WHERE pp.proveedor_id = ?
        ORDER BY p.nombre ASC
    """, (proveedor_id,))

    productos = cursor.fetchall()
    conexion.close()

    return jsonify([dict(producto) for producto in productos])


@app.route("/api/proveedores/<int:proveedor_id>/productos", methods=["POST"])
def agregar_producto_a_proveedor(proveedor_id):
    datos = request.json or {}

    nombre = str(datos.get("nombre", "") or "").strip()

    if not nombre:
        return jsonify({"mensaje": "El nombre del producto es obligatorio"}), 400

    variantes = datos.get("variantes")
    if variantes is None:
        variantes = []

    if not isinstance(variantes, list):
        return jsonify({"mensaje": "Las variantes deben enviarse como una lista"}), 400

    precio_base = float(datos.get("precio_base", 0) or 0)
    bonif1 = float(datos.get("bonificacion_1", 0) or 0)
    bonif2 = float(datos.get("bonificacion_2", 0) or 0)
    bonif3 = float(datos.get("bonificacion_3", 0) or 0)
    iva = float(datos.get("iva", 21) or 21)
    flete = float(datos.get("flete", 0) or 0)
    ganancia = float(datos.get("ganancia", 0) or 0)

    precio_final_compra, precio_venta_calculado = calcular_precios(
        precio_base, bonif1, bonif2, bonif3, iva, flete, ganancia
    )

    # Mientras adaptamos todo el frontend, los campos del producto padre
    # reflejan la primera variante. Así las pantallas antiguas siguen funcionando.
    if variantes:
        primera = variantes[0] or {}
        codigo_padre = str(primera.get("codigo", "") or "").strip()
        cantidad_caja_padre = int(primera.get("cantidad_caja", 1) or 1)
        precio_unidad_padre = float(primera.get("precio_unidad", 0) or 0)
        precio_caja_padre = float(primera.get("precio_caja", 0) or 0)
        precio_venta_padre = float(
            primera.get("precio_venta", precio_venta_calculado) or precio_venta_calculado or 0
        )
        stock_padre = sum(int((variante or {}).get("stock", 0) or 0) for variante in variantes)
    else:
        codigo_padre = str(datos.get("codigo", "") or "").strip()
        cantidad_caja_padre = int(datos.get("cantidad_caja", 1) or 1)
        precio_unidad_padre = float(datos.get("precio_unidad", 0) or 0)
        precio_caja_padre = float(datos.get("precio_caja", 0) or 0)
        precio_venta_padre = float(
            datos.get("precio_venta", precio_venta_calculado) or precio_venta_calculado or 0
        )
        stock_padre = int(datos.get("stock", 0) or 0)

    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        cursor.execute("SELECT id FROM proveedores WHERE id = ?", (proveedor_id,))
        if not cursor.fetchone():
            conexion.close()
            return jsonify({"mensaje": "Proveedor no encontrado"}), 404

        cursor.execute("""
            INSERT INTO productos (
                proveedor_id, codigo, nombre, descripcion, categoria, foto,
                cantidad_caja, precio_unidad, precio_caja, precio_venta, stock,
                foto_origen, foto_palabra_clave
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            proveedor_id,
            codigo_padre,
            nombre,
            datos.get("descripcion", ""),
            datos.get("categoria", ""),
            datos.get("foto", ""),
            cantidad_caja_padre,
            precio_unidad_padre,
            precio_caja_padre,
            precio_venta_padre,
            stock_padre,
            datos.get("foto_origen", ""),
            datos.get("foto_palabra_clave", "")
        ))

        producto_id = cursor.lastrowid

        for variante in variantes:
            variante = variante or {}
            nombre_variante = str(variante.get("nombre_variante", "") or "").strip()

            if not nombre_variante:
                raise ValueError("Todas las variantes deben tener una medida, modelo o nombre")

            cursor.execute("""
                INSERT INTO producto_variantes (
                    producto_id, nombre_variante, codigo,
                    cantidad_caja, precio_unidad, precio_caja,
                    precio_venta, stock, activo
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                producto_id,
                nombre_variante,
                str(variante.get("codigo", "") or "").strip(),
                int(variante.get("cantidad_caja", 1) or 1),
                float(variante.get("precio_unidad", 0) or 0),
                float(variante.get("precio_caja", 0) or 0),
                float(variante.get("precio_venta", precio_venta_calculado) or precio_venta_calculado or 0),
                int(variante.get("stock", 0) or 0),
                1
            ))

        cursor.execute("""
            INSERT INTO precios_proveedor (
                proveedor_id, producto_id, precio_base,
                bonificacion_1, bonificacion_2, bonificacion_3,
                iva, flete, precio_final_compra, ganancia,
                precio_venta, fecha_actualizacion
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            proveedor_id,
            producto_id,
            precio_base,
            bonif1,
            bonif2,
            bonif3,
            iva,
            flete,
            precio_final_compra,
            ganancia,
            precio_venta_padre,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ))

        conexion.commit()

        return jsonify({
            "mensaje": "Producto agregado al proveedor correctamente",
            "producto_id": producto_id,
            "cantidad_variantes": len(variantes)
        }), 201

    except (ValueError, TypeError) as error:
        conexion.rollback()
        return jsonify({"mensaje": str(error)}), 400
    except Exception as error:
        conexion.rollback()
        print("ERROR CREAR PRODUCTO:", error)
        return jsonify({"mensaje": "No se pudo guardar el producto", "error": str(error)}), 500
    finally:
        conexion.close()


@app.route("/api/precios/<int:precio_id>", methods=["PUT"])
def actualizar_precio_proveedor(precio_id):
    datos = request.json or {}

    precio_base = float(datos.get("precio_base", 0) or 0)
    bonif1 = float(datos.get("bonificacion_1", 0) or 0)
    bonif2 = float(datos.get("bonificacion_2", 0) or 0)
    bonif3 = float(datos.get("bonificacion_3", 0) or 0)
    iva = float(datos.get("iva", 21) or 21)
    flete = float(datos.get("flete", 0) or 0)
    ganancia = float(datos.get("ganancia", 0) or 0)

    precio_final_compra, precio_venta = calcular_precios(
        precio_base, bonif1, bonif2, bonif3, iva, flete, ganancia
    )

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        UPDATE precios_proveedor
        SET precio_base = ?,
            bonificacion_1 = ?,
            bonificacion_2 = ?,
            bonificacion_3 = ?,
            iva = ?,
            flete = ?,
            precio_final_compra = ?,
            ganancia = ?,
            precio_venta = ?,
            fecha_actualizacion = ?
        WHERE id = ?
    """, (
        precio_base,
        bonif1,
        bonif2,
        bonif3,
        iva,
        flete,
        precio_final_compra,
        ganancia,
        precio_venta,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        precio_id
    ))

    cursor.execute("""
        UPDATE productos
        SET precio_venta = ?
        WHERE id = (
            SELECT producto_id FROM precios_proveedor WHERE id = ?
        )
    """, (precio_venta, precio_id))

    conexion.commit()
    conexion.close()

    return jsonify({
        "mensaje": "Precio actualizado correctamente",
        "precio_final_compra": precio_final_compra,
        "precio_venta": precio_venta
    })


@app.route("/api/productos", methods=["GET"])
def listar_productos():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            p.*,
            pr.nombre AS proveedor
        FROM productos p
        LEFT JOIN proveedores pr ON p.proveedor_id = pr.id
        ORDER BY p.id DESC
    """)

    productos = cursor.fetchall()
    lista = []

    for producto in productos:
        cursor.execute("""
            SELECT
                id,
                producto_id,
                nombre_variante,
                codigo,
                cantidad_caja,
                precio_unidad,
                precio_caja,
                precio_venta,
                stock,
                activo
            FROM producto_variantes
            WHERE producto_id = ?
              AND activo = 1
            ORDER BY id ASC
        """, (producto["id"],))

        variantes = [dict(variante) for variante in cursor.fetchall()]

        lista.append({
            "id": producto["id"],
            "proveedor_id": producto["proveedor_id"],
            "proveedor": producto["proveedor"],
            "codigo": producto["codigo"],
            "nombre": producto["nombre"],
            "descripcion": producto["descripcion"],
            "categoria": producto["categoria"],
            "foto": producto["foto"],
            "foto_origen": producto["foto_origen"] if "foto_origen" in producto.keys() else "",
            "foto_palabra_clave": producto["foto_palabra_clave"] if "foto_palabra_clave" in producto.keys() else "",
            "cantidad_caja": producto["cantidad_caja"],
            "precio_unidad": producto["precio_unidad"],
            "precio_caja": producto["precio_caja"],
            "precio": producto["precio_venta"],
            "precio_venta": producto["precio_venta"],
            "stock": producto["stock"],
            "variantes": variantes,
            "cantidad_variantes": len(variantes)
        })

    conexion.close()
    return jsonify(lista)


@app.route("/api/productos/<int:id>", methods=["PUT"])
def editar_producto(id):
    datos = request.json or {}

    nombre = str(datos.get("nombre", "") or "").strip()
    if not nombre:
        return jsonify({"mensaje": "El nombre del producto es obligatorio"}), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        cursor.execute("SELECT * FROM productos WHERE id = ?", (id,))
        producto_actual = cursor.fetchone()

        if not producto_actual:
            conexion.close()
            return jsonify({"mensaje": "Producto no encontrado"}), 404

        tiene_variantes_en_payload = "variantes" in datos
        variantes = datos.get("variantes", []) if tiene_variantes_en_payload else None

        if tiene_variantes_en_payload and not isinstance(variantes, list):
            conexion.close()
            return jsonify({"mensaje": "Las variantes deben enviarse como una lista"}), 400

        if tiene_variantes_en_payload and variantes:
            primera = variantes[0] or {}
            codigo_padre = str(primera.get("codigo", "") or "").strip()
            cantidad_caja_padre = int(primera.get("cantidad_caja", 1) or 1)
            precio_unidad_padre = float(primera.get("precio_unidad", 0) or 0)
            precio_caja_padre = float(primera.get("precio_caja", 0) or 0)
            precio_venta_padre = float(primera.get("precio_venta", 0) or 0)
            stock_padre = sum(int((variante or {}).get("stock", 0) or 0) for variante in variantes)
        else:
            codigo_padre = datos.get("codigo", producto_actual["codigo"] or "")
            cantidad_caja_padre = int(datos.get("cantidad_caja", producto_actual["cantidad_caja"] or 1) or 1)
            precio_unidad_padre = float(datos.get("precio_unidad", producto_actual["precio_unidad"] or 0) or 0)
            precio_caja_padre = float(datos.get("precio_caja", producto_actual["precio_caja"] or 0) or 0)
            precio_venta_padre = float(
                datos.get("precio_venta", datos.get("precio", producto_actual["precio_venta"] or 0)) or 0
            )
            stock_padre = int(datos.get("stock", producto_actual["stock"] or 0) or 0)

        cursor.execute("""
            UPDATE productos
            SET codigo = ?,
                nombre = ?,
                descripcion = ?,
                categoria = ?,
                foto = ?,
                cantidad_caja = ?,
                precio_unidad = ?,
                precio_caja = ?,
                precio_venta = ?,
                stock = ?,
                foto_origen = ?,
                foto_palabra_clave = ?
            WHERE id = ?
        """, (
            codigo_padre,
            nombre,
            datos.get("descripcion", producto_actual["descripcion"] or ""),
            datos.get("categoria", producto_actual["categoria"] or ""),
            datos.get("foto", producto_actual["foto"] or ""),
            cantidad_caja_padre,
            precio_unidad_padre,
            precio_caja_padre,
            precio_venta_padre,
            stock_padre,
            datos.get("foto_origen", producto_actual["foto_origen"] if "foto_origen" in producto_actual.keys() else ""),
            datos.get("foto_palabra_clave", producto_actual["foto_palabra_clave"] if "foto_palabra_clave" in producto_actual.keys() else ""),
            id
        ))

        if tiene_variantes_en_payload:
            ids_recibidos = []

            for variante in variantes:
                variante = variante or {}
                nombre_variante = str(variante.get("nombre_variante", "") or "").strip()

                if not nombre_variante:
                    raise ValueError("Todas las variantes deben tener una medida, modelo o nombre")

                variante_id = variante.get("id")

                if variante_id:
                    cursor.execute("""
                        SELECT id
                        FROM producto_variantes
                        WHERE id = ? AND producto_id = ?
                    """, (variante_id, id))

                    existe = cursor.fetchone()

                    if existe:
                        cursor.execute("""
                            UPDATE producto_variantes
                            SET nombre_variante = ?,
                                codigo = ?,
                                cantidad_caja = ?,
                                precio_unidad = ?,
                                precio_caja = ?,
                                precio_venta = ?,
                                stock = ?,
                                activo = 1
                            WHERE id = ? AND producto_id = ?
                        """, (
                            nombre_variante,
                            str(variante.get("codigo", "") or "").strip(),
                            int(variante.get("cantidad_caja", 1) or 1),
                            float(variante.get("precio_unidad", 0) or 0),
                            float(variante.get("precio_caja", 0) or 0),
                            float(variante.get("precio_venta", 0) or 0),
                            int(variante.get("stock", 0) or 0),
                            variante_id,
                            id
                        ))
                        ids_recibidos.append(int(variante_id))
                        continue

                cursor.execute("""
                    INSERT INTO producto_variantes (
                        producto_id, nombre_variante, codigo,
                        cantidad_caja, precio_unidad, precio_caja,
                        precio_venta, stock, activo
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    id,
                    nombre_variante,
                    str(variante.get("codigo", "") or "").strip(),
                    int(variante.get("cantidad_caja", 1) or 1),
                    float(variante.get("precio_unidad", 0) or 0),
                    float(variante.get("precio_caja", 0) or 0),
                    float(variante.get("precio_venta", 0) or 0),
                    int(variante.get("stock", 0) or 0),
                    1
                ))
                ids_recibidos.append(cursor.lastrowid)

            if ids_recibidos:
                placeholders = ",".join("?" for _ in ids_recibidos)
                cursor.execute(
                    f"DELETE FROM producto_variantes WHERE producto_id = ? AND id NOT IN ({placeholders})",
                    [id] + ids_recibidos
                )
            else:
                cursor.execute("DELETE FROM producto_variantes WHERE producto_id = ?", (id,))

        conexion.commit()

        return jsonify({
            "mensaje": "Producto actualizado correctamente",
            "producto_id": id,
            "cantidad_variantes": len(variantes) if tiene_variantes_en_payload else None
        })

    except (ValueError, TypeError) as error:
        conexion.rollback()
        return jsonify({"mensaje": str(error)}), 400
    except Exception as error:
        conexion.rollback()
        print("ERROR EDITAR PRODUCTO:", error)
        return jsonify({"mensaje": "No se pudo actualizar el producto", "error": str(error)}), 500
    finally:
        conexion.close()


@app.route("/api/productos/<int:id>", methods=["DELETE"])
def eliminar_producto(id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        cursor.execute("SELECT id FROM productos WHERE id = ?", (id,))
        if not cursor.fetchone():
            conexion.close()
            return jsonify({"mensaje": "Producto no encontrado"}), 404

        cursor.execute("DELETE FROM producto_variantes WHERE producto_id = ?", (id,))
        cursor.execute("DELETE FROM precios_proveedor WHERE producto_id = ?", (id,))
        cursor.execute("DELETE FROM productos WHERE id = ?", (id,))

        conexion.commit()

        return jsonify({
            "mensaje": "Producto y sus variantes eliminados correctamente"
        })

    except Exception as error:
        conexion.rollback()
        print("ERROR ELIMINAR PRODUCTO:", error)
        return jsonify({"mensaje": "No se pudo eliminar el producto", "error": str(error)}), 500
    finally:
        conexion.close()


def _leer_porcentaje_aumento(datos):
    valor = str(datos.get("porcentaje", "") or "").strip().replace(",", ".")

    try:
        porcentaje = Decimal(valor)
    except (InvalidOperation, ValueError):
        raise ValueError("Ingresá un porcentaje válido")

    if not porcentaje.is_finite() or porcentaje <= 0:
        raise ValueError("El porcentaje debe ser mayor a cero")

    if porcentaje > Decimal("1000"):
        raise ValueError("El porcentaje no puede superar el 1000%")

    return float(porcentaje)


def _leer_ids_productos(datos):
    valores = datos.get("producto_ids") or []

    if not isinstance(valores, list):
        raise ValueError("La selección de productos no es válida")

    ids = []
    for valor in valores:
        try:
            producto_id = int(valor)
        except (TypeError, ValueError):
            raise ValueError("La selección contiene un producto inválido")

        if producto_id > 0 and producto_id not in ids:
            ids.append(producto_id)

    return ids


def _seleccionar_productos_para_aumento(cursor, datos):
    modo = str(datos.get("modo", "") or "").strip().lower()

    if modo not in ("general", "parcial", "individual"):
        raise ValueError("Seleccioná un tipo de aumento válido")

    producto_ids = _leer_ids_productos(datos)
    proveedor_id = datos.get("proveedor_id")
    categoria = str(datos.get("categoria", "") or "").strip()

    if proveedor_id in (None, ""):
        proveedor_id = None
    else:
        try:
            proveedor_id = int(proveedor_id)
        except (TypeError, ValueError):
            raise ValueError("El proveedor seleccionado no es válido")

    if modo == "individual" and len(producto_ids) != 1:
        raise ValueError("Seleccioná un producto para el aumento individual")

    if modo == "parcial" and not (producto_ids or proveedor_id or categoria):
        raise ValueError(
            "Elegí un proveedor, una categoría o productos para el aumento parcial"
        )

    condiciones = []
    parametros = []

    if modo == "individual":
        condiciones.append("p.id = ?")
        parametros.append(producto_ids[0])
    elif modo == "parcial":
        if proveedor_id:
            condiciones.append("p.proveedor_id = ?")
            parametros.append(proveedor_id)

        if categoria:
            condiciones.append("LOWER(TRIM(COALESCE(p.categoria, ''))) = LOWER(?)")
            parametros.append(categoria)

    clausula_where = ""
    if condiciones:
        clausula_where = "WHERE " + " AND ".join(condiciones)

    cursor.execute(f"""
        SELECT
            p.id,
            p.proveedor_id,
            p.codigo,
            p.nombre,
            p.categoria,
            p.precio_venta,
            pr.nombre AS proveedor
        FROM productos p
        LEFT JOIN proveedores pr ON pr.id = p.proveedor_id
        {clausula_where}
        ORDER BY p.nombre COLLATE NOCASE ASC, p.id ASC
    """, parametros)

    productos = [dict(fila) for fila in cursor.fetchall()]

    if modo == "parcial" and producto_ids:
        ids_permitidos = set(producto_ids)
        productos = [
            producto for producto in productos
            if producto["id"] in ids_permitidos
        ]

    if not productos:
        raise ValueError("No hay productos que coincidan con la selección")

    if modo == "general":
        criterio = "Todos los productos con precio de venta"
    elif modo == "individual":
        criterio = f"Producto: {productos[0]['nombre']}"
    else:
        partes = []

        if proveedor_id:
            cursor.execute(
                "SELECT nombre FROM proveedores WHERE id = ?",
                (proveedor_id,)
            )
            proveedor = cursor.fetchone()
            partes.append(
                f"Proveedor: {proveedor['nombre'] if proveedor else proveedor_id}"
            )

        if categoria:
            partes.append(f"Categoría: {categoria}")

        if producto_ids:
            partes.append(f"Selección manual: {len(productos)} producto(s)")

        criterio = " · ".join(partes)

    return modo, criterio, productos


def _construir_vista_previa_aumento(cursor, productos, porcentaje):
    productos_por_id = {producto["id"]: producto for producto in productos}
    variantes_por_producto = {}
    ids = list(productos_por_id.keys())

    for inicio in range(0, len(ids), 800):
        bloque = ids[inicio:inicio + 800]
        placeholders = ",".join("?" for _ in bloque)
        cursor.execute(f"""
            SELECT
                id,
                producto_id,
                nombre_variante,
                codigo,
                precio_venta
            FROM producto_variantes
            WHERE activo = 1
              AND producto_id IN ({placeholders})
            ORDER BY producto_id ASC, id ASC
        """, bloque)

        for fila in cursor.fetchall():
            variante = dict(fila)
            variantes_por_producto.setdefault(
                variante["producto_id"], []
            ).append(variante)

    vista = []
    omitidos_sin_precio = 0
    cantidad_precios = 0

    for producto in productos:
        lineas = []

        for variante in variantes_por_producto.get(producto["id"], []):
            precio_anterior = float(variante["precio_venta"] or 0)
            if precio_anterior <= 0:
                continue

            lineas.append({
                "variante_id": variante["id"],
                "variante": variante["nombre_variante"] or "Presentación",
                "codigo": variante["codigo"] or producto["codigo"] or "",
                "precio_anterior": precio_anterior,
                "precio_nuevo": redondear_aumento_hacia_arriba_10(
                    precio_anterior,
                    porcentaje
                )
            })

        if not lineas:
            precio_anterior = float(producto["precio_venta"] or 0)

            if precio_anterior > 0:
                lineas.append({
                    "variante_id": None,
                    "variante": "Precio principal",
                    "codigo": producto["codigo"] or "",
                    "precio_anterior": precio_anterior,
                    "precio_nuevo": redondear_aumento_hacia_arriba_10(
                        precio_anterior,
                        porcentaje
                    )
                })

        if not lineas:
            omitidos_sin_precio += 1
            continue

        cantidad_precios += len(lineas)
        vista.append({
            "id": producto["id"],
            "nombre": producto["nombre"],
            "codigo": producto["codigo"] or lineas[0]["codigo"],
            "proveedor": producto["proveedor"] or "Sin proveedor",
            "categoria": producto["categoria"] or "Sin categoría",
            "precio_anterior": lineas[0]["precio_anterior"],
            "precio_nuevo": lineas[0]["precio_nuevo"],
            "precios": lineas
        })

    if not vista:
        raise ValueError("Los productos seleccionados no tienen precio de venta")

    return {
        "productos": vista,
        "cantidad_productos": len(vista),
        "cantidad_precios": cantidad_precios,
        "omitidos_sin_precio": omitidos_sin_precio
    }


@app.route("/api/aumentos-precios/opciones", methods=["GET"])
def opciones_aumentos_precios():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            p.id,
            p.proveedor_id,
            p.codigo,
            p.nombre,
            p.categoria,
            COALESCE(
                (
                    SELECT pv.precio_venta
                    FROM producto_variantes pv
                    WHERE pv.producto_id = p.id
                      AND pv.activo = 1
                      AND pv.precio_venta > 0
                    ORDER BY pv.id ASC
                    LIMIT 1
                ),
                p.precio_venta,
                0
            ) AS precio_actual,
            (
                SELECT COUNT(*)
                FROM producto_variantes pv
                WHERE pv.producto_id = p.id AND pv.activo = 1
            ) AS cantidad_variantes,
            pr.nombre AS proveedor
        FROM productos p
        LEFT JOIN proveedores pr ON pr.id = p.proveedor_id
        ORDER BY p.nombre COLLATE NOCASE ASC, p.id ASC
    """)
    productos = [dict(fila) for fila in cursor.fetchall()]

    cursor.execute("""
        SELECT id, nombre
        FROM proveedores
        ORDER BY nombre COLLATE NOCASE ASC
    """)
    proveedores = [dict(fila) for fila in cursor.fetchall()]

    cursor.execute("""
        SELECT DISTINCT TRIM(categoria) AS categoria
        FROM productos
        WHERE TRIM(COALESCE(categoria, '')) <> ''
        ORDER BY categoria COLLATE NOCASE ASC
    """)
    categorias = [fila["categoria"] for fila in cursor.fetchall()]

    conexion.close()
    respuesta = jsonify({
        "productos": productos,
        "proveedores": proveedores,
        "categorias": categorias,
        "redondeo": "Siempre hacia arriba al próximo múltiplo de $10"
    })
    respuesta.headers["Cache-Control"] = "no-store"
    return respuesta


@app.route("/api/aumentos-precios/vista-previa", methods=["POST"])
def vista_previa_aumento_precios():
    datos = request.json or {}
    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        porcentaje = _leer_porcentaje_aumento(datos)
        modo, criterio, productos = _seleccionar_productos_para_aumento(
            cursor,
            datos
        )
        vista = _construir_vista_previa_aumento(
            cursor,
            productos,
            porcentaje
        )
    except ValueError as error:
        conexion.close()
        return jsonify({"mensaje": str(error)}), 400

    conexion.close()
    return jsonify({
        "modo": modo,
        "porcentaje": porcentaje,
        "criterio": criterio,
        "redondeo": "Hacia arriba a $10",
        **vista
    })


@app.route("/api/aumentos-precios", methods=["POST"])
def aplicar_aumento_precios():
    datos = request.json or {}
    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        cursor.execute("BEGIN IMMEDIATE")
        porcentaje = _leer_porcentaje_aumento(datos)
        modo, criterio, productos = _seleccionar_productos_para_aumento(
            cursor,
            datos
        )
        vista = _construir_vista_previa_aumento(
            cursor,
            productos,
            porcentaje
        )
        fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        cursor.execute("""
            INSERT INTO aumentos_precios (
                fecha, modo, porcentaje, criterio,
                cantidad_productos, cantidad_precios, redondeo
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (
            fecha,
            modo,
            porcentaje,
            criterio,
            vista["cantidad_productos"],
            vista["cantidad_precios"],
            "Hacia arriba a $10"
        ))
        aumento_id = cursor.lastrowid

        for producto in vista["productos"]:
            for precio in producto["precios"]:
                if precio["variante_id"] is not None:
                    cursor.execute("""
                        UPDATE producto_variantes
                        SET precio_venta = ?
                        WHERE id = ? AND producto_id = ?
                    """, (
                        precio["precio_nuevo"],
                        precio["variante_id"],
                        producto["id"]
                    ))

                cursor.execute("""
                    INSERT INTO aumento_precio_detalles (
                        aumento_id, producto_id, producto,
                        variante_id, variante, codigo,
                        proveedor, categoria, precio_anterior, precio_nuevo
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    aumento_id,
                    producto["id"],
                    producto["nombre"],
                    precio["variante_id"],
                    precio["variante"],
                    precio["codigo"],
                    producto["proveedor"],
                    producto["categoria"],
                    precio["precio_anterior"],
                    precio["precio_nuevo"]
                ))

            nuevo_precio_principal = producto["precios"][0]["precio_nuevo"]
            cursor.execute("""
                UPDATE productos
                SET precio_venta = ?
                WHERE id = ?
            """, (nuevo_precio_principal, producto["id"]))

            cursor.execute("""
                UPDATE precios_proveedor
                SET precio_venta = ?,
                    ganancia = CASE
                        WHEN COALESCE(precio_final_compra, 0) > 0
                        THEN ROUND(
                            ((? / precio_final_compra) - 1) * 100,
                            2
                        )
                        ELSE ganancia
                    END,
                    fecha_actualizacion = ?
                WHERE producto_id = ?
            """, (
                nuevo_precio_principal,
                nuevo_precio_principal,
                fecha,
                producto["id"]
            ))

        conexion.commit()
    except ValueError as error:
        conexion.rollback()
        conexion.close()
        return jsonify({"mensaje": str(error)}), 400
    except Exception as error:
        conexion.rollback()
        conexion.close()
        print("ERROR AUMENTO PRECIOS:", error)
        return jsonify({
            "mensaje": "No se pudo aplicar el aumento",
            "error": str(error)
        }), 500

    conexion.close()
    return jsonify({
        "mensaje": "Aumento aplicado correctamente",
        "aumento_id": aumento_id,
        "modo": modo,
        "porcentaje": porcentaje,
        "cantidad_productos": vista["cantidad_productos"],
        "cantidad_precios": vista["cantidad_precios"],
        "omitidos_sin_precio": vista["omitidos_sin_precio"],
        "redondeo": "Hacia arriba a $10"
    }), 201


@app.route("/api/aumentos-precios/historial", methods=["GET"])
def historial_aumentos_precios():
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT
            id, fecha, modo, porcentaje, criterio,
            cantidad_productos, cantidad_precios, redondeo
        FROM aumentos_precios
        ORDER BY id DESC
        LIMIT 100
    """)
    historial = [dict(fila) for fila in cursor.fetchall()]
    conexion.close()
    return jsonify(historial)


@app.route("/api/aumentos-precios/historial/<int:aumento_id>", methods=["GET"])
def detalle_aumento_precios(aumento_id):
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT
            id, fecha, modo, porcentaje, criterio,
            cantidad_productos, cantidad_precios, redondeo
        FROM aumentos_precios
        WHERE id = ?
    """, (aumento_id,))
    aumento = cursor.fetchone()

    if not aumento:
        conexion.close()
        return jsonify({"mensaje": "Aumento no encontrado"}), 404

    cursor.execute("""
        SELECT
            id, producto_id, producto, variante_id, variante,
            codigo, proveedor, categoria, precio_anterior, precio_nuevo
        FROM aumento_precio_detalles
        WHERE aumento_id = ?
        ORDER BY producto COLLATE NOCASE ASC, id ASC
    """, (aumento_id,))
    detalles = [dict(fila) for fila in cursor.fetchall()]
    conexion.close()

    respuesta = dict(aumento)
    respuesta["detalles"] = detalles
    return jsonify(respuesta)


@app.route("/api/clientes", methods=["GET"])
def listar_clientes():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("SELECT * FROM clientes ORDER BY nombre ASC")

    clientes = cursor.fetchall()
    conexion.close()

    return jsonify([dict(cliente) for cliente in clientes])


@app.route("/api/clientes", methods=["POST"])
def agregar_cliente():
    datos = request.json or {}

    nombre = datos.get("nombre", "").strip()

    if not nombre:
        return jsonify({"mensaje": "El nombre del cliente es obligatorio"}), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    numero_cliente = datos.get("numero_cliente") or generar_numero_cliente(cursor)

    cursor.execute("""
        INSERT INTO clientes (
            numero_cliente, nombre, direccion, localidad, telefono, email, observaciones, saldo_actual
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        numero_cliente,
        nombre,
        datos.get("direccion", ""),
        datos.get("localidad", ""),
        datos.get("telefono", ""),
        datos.get("email", ""),
        datos.get("observaciones", ""),
        float(datos.get("saldo_actual", 0) or 0)
    ))

    cliente_id = cursor.lastrowid

    conexion.commit()
    conexion.close()

    return jsonify({
        "mensaje": "Cliente creado correctamente",
        "cliente_id": cliente_id,
        "numero_cliente": numero_cliente
    }), 201


@app.route("/api/clientes/<int:id>", methods=["PUT"])
def editar_cliente(id):
    datos = request.json or {}

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        UPDATE clientes
        SET numero_cliente = ?,
            nombre = ?,
            direccion = ?,
            localidad = ?,
            telefono = ?,
            email = ?,
            observaciones = ?,
            saldo_actual = ?
        WHERE id = ?
    """, (
        datos.get("numero_cliente"),
        datos.get("nombre"),
        datos.get("direccion", ""),
        datos.get("localidad", ""),
        datos.get("telefono", ""),
        datos.get("email", ""),
        datos.get("observaciones", ""),
        float(datos.get("saldo_actual", 0) or 0),
        id
    ))

    conexion.commit()
    conexion.close()

    return jsonify({
        "mensaje": "Cliente actualizado correctamente"
    })


@app.route("/api/clientes/<int:id>", methods=["DELETE"])
def eliminar_cliente(id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("DELETE FROM cuentas_corrientes WHERE cliente_id = ?", (id,))
    cursor.execute("DELETE FROM pagos WHERE cliente_id = ?", (id,))
    cursor.execute("DELETE FROM saldos WHERE cliente_id = ?", (id,))
    cursor.execute("DELETE FROM clientes WHERE id = ?", (id,))

    conexion.commit()
    conexion.close()

    return jsonify({
        "mensaje": "Cliente eliminado correctamente"
    })


@app.route("/api/clientes/<int:cliente_id>/cuenta-corriente", methods=["GET"])
def obtener_cuenta_corriente(cliente_id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("SELECT * FROM clientes WHERE id = ?", (cliente_id,))
    cliente = cursor.fetchone()

    if not cliente:
        conexion.close()
        return jsonify({
            "mensaje": "Cliente no encontrado"
        }), 404

    cursor.execute("""
        SELECT *
        FROM cuentas_corrientes
        WHERE cliente_id = ?
        ORDER BY id DESC
    """, (cliente_id,))

    movimientos = cursor.fetchall()

    conexion.close()

    return jsonify({
        "cliente": dict(cliente),
        "saldo_actual": cliente["saldo_actual"],
        "movimientos": [dict(movimiento) for movimiento in movimientos]
    })


@app.route("/api/clientes/<int:cliente_id>/cuenta-corriente/movimiento", methods=["POST"])
def agregar_movimiento_cuenta_corriente(cliente_id):
    datos = request.json or {}

    descripcion = datos.get("descripcion", "").strip()
    monto = float(datos.get("monto", 0) or 0)
    tipo = datos.get("tipo", "Venta").strip()

    if not descripcion:
        return jsonify({
            "mensaje": "La descripción es obligatoria"
        }), 400

    if monto <= 0:
        return jsonify({
            "mensaje": "El monto debe ser mayor a cero"
        }), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("SELECT * FROM clientes WHERE id = ?", (cliente_id,))
    cliente = cursor.fetchone()

    if not cliente:
        conexion.close()
        return jsonify({
            "mensaje": "Cliente no encontrado"
        }), 404

    saldo_anterior = float(cliente["saldo_actual"] or 0)
    nuevo_saldo = saldo_anterior + monto

    cursor.execute("""
        INSERT INTO cuentas_corrientes (
            cliente_id, fecha, tipo, descripcion, comprobante,
            debe, haber, saldo, medio_pago
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        cliente_id,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        tipo,
        descripcion,
        datos.get("comprobante", ""),
        monto,
        0,
        nuevo_saldo,
        ""
    ))

    cursor.execute("""
        UPDATE clientes
        SET saldo_actual = ?
        WHERE id = ?
    """, (nuevo_saldo, cliente_id))

    conexion.commit()
    conexion.close()

    return jsonify({
        "mensaje": "Movimiento agregado correctamente",
        "saldo_actual": nuevo_saldo
    }), 201


@app.route("/api/clientes/<int:cliente_id>/pagos", methods=["POST"])
def registrar_pago_cliente(cliente_id):
    datos = request.json or {}

    monto = float(datos.get("monto", 0) or 0)
    medio_pago = datos.get("medio_pago", "").strip()
    comprobante = datos.get("comprobante", "").strip()
    observaciones = datos.get("observaciones", "").strip()

    if monto <= 0:
        return jsonify({
            "mensaje": "El monto del pago debe ser mayor a cero"
        }), 400

    if not medio_pago:
        return jsonify({
            "mensaje": "El medio de pago es obligatorio"
        }), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("SELECT * FROM clientes WHERE id = ?", (cliente_id,))
    cliente = cursor.fetchone()

    if not cliente:
        conexion.close()
        return jsonify({
            "mensaje": "Cliente no encontrado"
        }), 404

    saldo_anterior = float(cliente["saldo_actual"] or 0)
    nuevo_saldo = saldo_anterior - monto

    if nuevo_saldo < 0:
        nuevo_saldo = 0

    fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cursor.execute("""
        INSERT INTO pagos (
            cliente_id, fecha, monto, medio_pago, comprobante, observaciones
        )
        VALUES (?, ?, ?, ?, ?, ?)
    """, (
        cliente_id,
        fecha,
        monto,
        medio_pago,
        comprobante,
        observaciones
    ))

    cursor.execute("""
        INSERT INTO cuentas_corrientes (
            cliente_id, fecha, tipo, descripcion, comprobante,
            debe, haber, saldo, medio_pago
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        cliente_id,
        fecha,
        "Pago",
        observaciones if observaciones else "Pago registrado",
        comprobante,
        0,
        monto,
        nuevo_saldo,
        medio_pago
    ))

    cursor.execute("""
        UPDATE clientes
        SET saldo_actual = ?
        WHERE id = ?
    """, (nuevo_saldo, cliente_id))

    conexion.commit()
    conexion.close()

    return jsonify({
        "mensaje": "Pago registrado correctamente",
        "saldo_actual": nuevo_saldo
    }), 201



@app.route("/api/vendedores", methods=["GET"])
def listar_vendedores():
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            v.id,
            v.usuario_id,
            v.nombre,
            v.telefono,
            v.zona,
            v.observaciones,
            v.activo,
            v.fecha_alta,
            u.usuario
        FROM vendedores v
        INNER JOIN usuarios u ON u.id = v.usuario_id
        ORDER BY v.activo DESC, v.nombre ASC
    """)

    vendedores = [dict(fila) for fila in cursor.fetchall()]
    conexion.close()
    return jsonify(vendedores)


@app.route("/api/vendedores", methods=["POST"])
def crear_vendedor():
    datos = request.json or {}

    nombre = limpiar_texto(datos.get("nombre"))
    usuario = limpiar_texto(datos.get("usuario"))
    contrasena = limpiar_texto(datos.get("contrasena"))

    if not nombre:
        return jsonify({"mensaje": "El nombre del vendedor es obligatorio"}), 400

    if not usuario:
        return jsonify({"mensaje": "El usuario del vendedor es obligatorio"}), 400

    if len(contrasena) < 4:
        return jsonify({"mensaje": "La contraseña debe tener al menos 4 caracteres"}), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        cursor.execute("""
            INSERT INTO usuarios (usuario, contrasena, rol, activo)
            VALUES (?, ?, 'vendedor', 1)
        """, (
            usuario,
            generate_password_hash(contrasena)
        ))
        usuario_id = cursor.lastrowid

        fecha_alta = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute("""
            INSERT INTO vendedores (
                usuario_id, nombre, telefono, zona,
                observaciones, activo, fecha_alta
            )
            VALUES (?, ?, ?, ?, ?, 1, ?)
        """, (
            usuario_id,
            nombre,
            limpiar_texto(datos.get("telefono")),
            limpiar_texto(datos.get("zona")),
            limpiar_texto(datos.get("observaciones")),
            fecha_alta
        ))

        vendedor_id = cursor.lastrowid
        conexion.commit()
    except sqlite3.IntegrityError:
        conexion.rollback()
        conexion.close()
        return jsonify({
            "mensaje": "Ese nombre de usuario ya está en uso"
        }), 409

    conexion.close()
    return jsonify({
        "mensaje": "Vendedor creado correctamente",
        "id": vendedor_id,
        "usuario_id": usuario_id
    }), 201


@app.route("/api/vendedores/<int:vendedor_id>", methods=["PUT"])
def actualizar_vendedor(vendedor_id):
    datos = request.json or {}
    nombre = limpiar_texto(datos.get("nombre"))
    usuario = limpiar_texto(datos.get("usuario"))
    contrasena = limpiar_texto(datos.get("contrasena"))
    activo = 1 if datos.get("activo", True) else 0

    if not nombre or not usuario:
        return jsonify({
            "mensaje": "Nombre y usuario son obligatorios"
        }), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT usuario_id
        FROM vendedores
        WHERE id = ?
    """, (vendedor_id,))
    vendedor = cursor.fetchone()

    if not vendedor:
        conexion.close()
        return jsonify({"mensaje": "Vendedor no encontrado"}), 404

    try:
        cursor.execute("""
            UPDATE vendedores
            SET nombre = ?, telefono = ?, zona = ?,
                observaciones = ?, activo = ?
            WHERE id = ?
        """, (
            nombre,
            limpiar_texto(datos.get("telefono")),
            limpiar_texto(datos.get("zona")),
            limpiar_texto(datos.get("observaciones")),
            activo,
            vendedor_id
        ))

        if contrasena:
            cursor.execute("""
                UPDATE usuarios
                SET usuario = ?, contrasena = ?, activo = ?
                WHERE id = ?
            """, (
                usuario,
                generate_password_hash(contrasena),
                activo,
                vendedor["usuario_id"]
            ))
        else:
            cursor.execute("""
                UPDATE usuarios
                SET usuario = ?, activo = ?
                WHERE id = ?
            """, (
                usuario,
                activo,
                vendedor["usuario_id"]
            ))

        conexion.commit()
    except sqlite3.IntegrityError:
        conexion.rollback()
        conexion.close()
        return jsonify({
            "mensaje": "Ese nombre de usuario ya está en uso"
        }), 409

    conexion.close()
    return jsonify({"mensaje": "Vendedor actualizado correctamente"})


@app.route("/api/vendedores/<int:vendedor_id>", methods=["DELETE"])
def desactivar_vendedor(vendedor_id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT usuario_id
        FROM vendedores
        WHERE id = ?
    """, (vendedor_id,))
    vendedor = cursor.fetchone()

    if not vendedor:
        conexion.close()
        return jsonify({"mensaje": "Vendedor no encontrado"}), 404

    cursor.execute("""
        UPDATE vendedores SET activo = 0 WHERE id = ?
    """, (vendedor_id,))
    cursor.execute("""
        UPDATE usuarios SET activo = 0 WHERE id = ?
    """, (vendedor["usuario_id"],))

    conexion.commit()
    conexion.close()

    return jsonify({"mensaje": "Vendedor desactivado correctamente"})


@app.route("/api/vendedores/por-usuario/<int:usuario_id>", methods=["GET"])
def vendedor_por_usuario(usuario_id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            v.id, v.nombre, v.telefono, v.zona,
            v.observaciones, v.activo, u.usuario
        FROM vendedores v
        INNER JOIN usuarios u ON u.id = v.usuario_id
        WHERE v.usuario_id = ?
    """, (usuario_id,))

    vendedor = cursor.fetchone()
    conexion.close()

    if not vendedor:
        return jsonify({"mensaje": "Vendedor no encontrado"}), 404

    return jsonify(dict(vendedor))


@app.route("/api/vendedores/<int:vendedor_id>/stock-viaje", methods=["GET"])
def obtener_stock_viaje(vendedor_id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT id FROM vendedores
        WHERE id = ? AND activo = 1
    """, (vendedor_id,))
    if not cursor.fetchone():
        conexion.close()
        return jsonify({"mensaje": "Vendedor no encontrado o inactivo"}), 404

    cursor.execute("""
        SELECT
            sv.id,
            sv.vendedor_id,
            sv.producto_id,
            sv.cantidad,
            sv.actualizado_en,
            p.codigo,
            p.nombre,
            p.descripcion,
            p.precio_venta,
            p.stock AS stock_central,
            pr.nombre AS proveedor
        FROM stock_viaje sv
        INNER JOIN productos p ON p.id = sv.producto_id
        LEFT JOIN proveedores pr ON pr.id = p.proveedor_id
        WHERE sv.vendedor_id = ?
          AND sv.cantidad > 0
        ORDER BY p.nombre ASC
    """, (vendedor_id,))

    stock = [dict(fila) for fila in cursor.fetchall()]
    conexion.close()
    return jsonify(stock)


@app.route("/api/vendedores/<int:vendedor_id>/catalogo-visual", methods=["GET"])
def obtener_catalogo_visual(vendedor_id):
    """Catálogo para mostrar al cliente durante la venta ambulante.

    La respuesta excluye precios, costos, proveedores y cantidades internas.
    Solo informa si cada producto está disponible en el stock móvil del
    vendedor indicado.
    """
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT id, nombre, activo
        FROM vendedores
        WHERE id = ?
    """, (vendedor_id,))
    vendedor = cursor.fetchone()

    if not vendedor or not vendedor["activo"]:
        conexion.close()
        return jsonify({
            "mensaje": "Vendedor no encontrado o inactivo"
        }), 404

    cursor.execute("""
        SELECT
            p.id,
            p.codigo,
            p.nombre,
            p.descripcion,
            p.categoria,
            p.foto,
            COALESCE(sv.cantidad, 0) AS cantidad_viaje
        FROM productos p
        LEFT JOIN stock_viaje sv
          ON sv.producto_id = p.id
         AND sv.vendedor_id = ?
        ORDER BY p.nombre COLLATE NOCASE ASC, p.id ASC
    """, (vendedor_id,))
    productos = cursor.fetchall()

    cursor.execute("""
        SELECT
            id,
            producto_id,
            nombre_variante,
            codigo
        FROM producto_variantes
        WHERE activo = 1
        ORDER BY producto_id ASC, nombre_variante COLLATE NOCASE ASC, id ASC
    """)

    variantes_por_producto = {}
    for variante in cursor.fetchall():
        producto_id = variante["producto_id"]
        variantes_por_producto.setdefault(producto_id, []).append({
            "id": variante["id"],
            "nombre_variante": variante["nombre_variante"],
            "codigo": variante["codigo"]
        })

    catalogo = []
    for producto in productos:
        catalogo.append({
            "id": producto["id"],
            "codigo": producto["codigo"] or "",
            "nombre": producto["nombre"],
            "descripcion": producto["descripcion"] or "",
            "categoria": producto["categoria"] or "Sin categoría",
            "foto": producto["foto"] or "",
            "disponible": int(producto["cantidad_viaje"] or 0) > 0,
            "variantes": variantes_por_producto.get(producto["id"], [])
        })

    conexion.close()

    respuesta = jsonify({
        "vendedor_id": vendedor["id"],
        "vendedor": vendedor["nombre"],
        "actualizado_en": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "productos": catalogo
    })
    respuesta.headers["Cache-Control"] = "no-store"
    return respuesta


@app.route("/api/vendedores/<int:vendedor_id>/stock-viaje/asignar", methods=["POST"])
def asignar_stock_viaje(vendedor_id):
    datos = request.json or {}

    try:
        producto_id = int(datos.get("producto_id"))
        cantidad = int(datos.get("cantidad"))
    except (TypeError, ValueError):
        return jsonify({
            "mensaje": "Producto y cantidad son obligatorios"
        }), 400

    if cantidad <= 0:
        return jsonify({"mensaje": "La cantidad debe ser mayor a cero"}), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        cursor.execute("""
            SELECT id FROM vendedores
            WHERE id = ? AND activo = 1
        """, (vendedor_id,))
        if not cursor.fetchone():
            conexion.close()
            return jsonify({"mensaje": "Vendedor no encontrado o inactivo"}), 404

        cursor.execute("""
            SELECT id, nombre, stock
            FROM productos
            WHERE id = ?
        """, (producto_id,))
        producto = cursor.fetchone()

        if not producto:
            conexion.close()
            return jsonify({"mensaje": "Producto no encontrado"}), 404

        stock_central = int(producto["stock"] or 0)

        if stock_central < cantidad:
            conexion.close()
            return jsonify({
                "mensaje": (
                    f"Stock central insuficiente. Disponible: {stock_central}"
                )
            }), 400

        fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        nuevo_stock_central = stock_central - cantidad

        cursor.execute("""
            SELECT id, cantidad
            FROM stock_viaje
            WHERE vendedor_id = ? AND producto_id = ?
        """, (vendedor_id, producto_id))
        registro = cursor.fetchone()

        if registro:
            nuevo_stock_viaje = int(registro["cantidad"] or 0) + cantidad
            cursor.execute("""
                UPDATE stock_viaje
                SET cantidad = ?, actualizado_en = ?
                WHERE id = ?
            """, (
                nuevo_stock_viaje,
                fecha,
                registro["id"]
            ))
        else:
            nuevo_stock_viaje = cantidad
            cursor.execute("""
                INSERT INTO stock_viaje (
                    vendedor_id, producto_id, cantidad, actualizado_en
                )
                VALUES (?, ?, ?, ?)
            """, (
                vendedor_id,
                producto_id,
                cantidad,
                fecha
            ))

        cursor.execute("""
            UPDATE productos
            SET stock = ?
            WHERE id = ?
        """, (nuevo_stock_central, producto_id))

        cursor.execute("""
            INSERT INTO movimientos_stock_viaje (
                vendedor_id, producto_id, fecha, tipo, cantidad,
                descripcion, stock_viaje_resultante, stock_central_resultante
            )
            VALUES (?, ?, ?, 'Carga', ?, ?, ?, ?)
        """, (
            vendedor_id,
            producto_id,
            fecha,
            cantidad,
            limpiar_texto(datos.get("descripcion")) or "Carga de mercadería",
            nuevo_stock_viaje,
            nuevo_stock_central
        ))

        cursor.execute("""
            INSERT INTO movimientos_stock (
                producto_id, fecha, tipo, cantidad, descripcion, stock_resultante
            )
            VALUES (?, ?, 'Salida a viaje', ?, ?, ?)
        """, (
            producto_id,
            fecha,
            -cantidad,
            f"Mercadería asignada al vendedor #{vendedor_id}",
            nuevo_stock_central
        ))

        conexion.commit()
    except Exception:
        conexion.rollback()
        conexion.close()
        raise

    conexion.close()
    return jsonify({
        "mensaje": "Mercadería asignada al stock en viaje",
        "stock_viaje": nuevo_stock_viaje,
        "stock_central": nuevo_stock_central
    }), 201


@app.route("/api/vendedores/<int:vendedor_id>/stock-viaje/devolver", methods=["POST"])
def devolver_stock_viaje(vendedor_id):
    datos = request.json or {}

    try:
        producto_id = int(datos.get("producto_id"))
        cantidad = int(datos.get("cantidad"))
    except (TypeError, ValueError):
        return jsonify({
            "mensaje": "Producto y cantidad son obligatorios"
        }), 400

    if cantidad <= 0:
        return jsonify({"mensaje": "La cantidad debe ser mayor a cero"}), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT cantidad
        FROM stock_viaje
        WHERE vendedor_id = ? AND producto_id = ?
    """, (vendedor_id, producto_id))
    registro = cursor.fetchone()

    if not registro or int(registro["cantidad"] or 0) < cantidad:
        conexion.close()
        return jsonify({
            "mensaje": "El vendedor no tiene esa cantidad para devolver"
        }), 400

    cursor.execute("""
        SELECT stock
        FROM productos
        WHERE id = ?
    """, (producto_id,))
    producto = cursor.fetchone()

    if not producto:
        conexion.close()
        return jsonify({"mensaje": "Producto no encontrado"}), 404

    stock_viaje_actual = int(registro["cantidad"] or 0)
    stock_central_actual = int(producto["stock"] or 0)
    nuevo_stock_viaje = stock_viaje_actual - cantidad
    nuevo_stock_central = stock_central_actual + cantidad
    fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cursor.execute("""
        UPDATE stock_viaje
        SET cantidad = ?, actualizado_en = ?
        WHERE vendedor_id = ? AND producto_id = ?
    """, (
        nuevo_stock_viaje,
        fecha,
        vendedor_id,
        producto_id
    ))

    cursor.execute("""
        UPDATE productos
        SET stock = ?
        WHERE id = ?
    """, (nuevo_stock_central, producto_id))

    cursor.execute("""
        INSERT INTO movimientos_stock_viaje (
            vendedor_id, producto_id, fecha, tipo, cantidad,
            descripcion, stock_viaje_resultante, stock_central_resultante
        )
        VALUES (?, ?, ?, 'Devolución', ?, ?, ?, ?)
    """, (
        vendedor_id,
        producto_id,
        fecha,
        -cantidad,
        limpiar_texto(datos.get("descripcion")) or "Devolución a depósito",
        nuevo_stock_viaje,
        nuevo_stock_central
    ))

    cursor.execute("""
        INSERT INTO movimientos_stock (
            producto_id, fecha, tipo, cantidad, descripcion, stock_resultante
        )
        VALUES (?, ?, 'Retorno de viaje', ?, ?, ?)
    """, (
        producto_id,
        fecha,
        cantidad,
        f"Devolución del vendedor #{vendedor_id}",
        nuevo_stock_central
    ))

    conexion.commit()
    conexion.close()

    return jsonify({
        "mensaje": "Mercadería devuelta correctamente",
        "stock_viaje": nuevo_stock_viaje,
        "stock_central": nuevo_stock_central
    })


@app.route("/api/vendedores/<int:vendedor_id>/stock-viaje/movimientos", methods=["GET"])
def listar_movimientos_stock_viaje(vendedor_id):
    conexion = conectar_db()
    cursor = conexion.cursor()

    cursor.execute("""
        SELECT
            m.*,
            p.codigo,
            p.nombre AS producto
        FROM movimientos_stock_viaje m
        INNER JOIN productos p ON p.id = m.producto_id
        WHERE m.vendedor_id = ?
        ORDER BY m.id DESC
        LIMIT 200
    """, (vendedor_id,))

    movimientos = [dict(fila) for fila in cursor.fetchall()]
    conexion.close()
    return jsonify(movimientos)


def _consultar_ventas(cursor, vendedor_id=None):
    parametros = []
    condicion = ""

    if vendedor_id is not None:
        condicion = "WHERE ve.vendedor_id = ?"
        parametros.append(vendedor_id)

    cursor.execute(f"""
        SELECT
            ve.id,
            ve.vendedor_id,
            ve.cliente_id,
            ve.fecha,
            ve.forma_pago,
            ve.total,
            ve.observaciones,
            v.nombre AS vendedor,
            c.nombre AS cliente,
            c.numero_cliente
        FROM ventas ve
        INNER JOIN vendedores v ON v.id = ve.vendedor_id
        LEFT JOIN clientes c ON c.id = ve.cliente_id
        {condicion}
        ORDER BY ve.id DESC
        LIMIT 300
    """, parametros)

    ventas = [dict(fila) for fila in cursor.fetchall()]

    for venta in ventas:
        cursor.execute("""
            SELECT
                vd.id,
                vd.producto_id,
                vd.cantidad,
                vd.precio_unitario,
                vd.subtotal,
                p.codigo,
                p.nombre AS producto
            FROM venta_detalles vd
            INNER JOIN productos p ON p.id = vd.producto_id
            WHERE vd.venta_id = ?
            ORDER BY vd.id ASC
        """, (venta["id"],))
        venta["detalles"] = [dict(fila) for fila in cursor.fetchall()]

    return ventas


@app.route("/api/ventas", methods=["GET"])
def listar_ventas():
    conexion = conectar_db()
    cursor = conexion.cursor()
    ventas = _consultar_ventas(cursor)
    conexion.close()
    return jsonify(ventas)


@app.route("/api/vendedores/<int:vendedor_id>/ventas", methods=["GET"])
def listar_ventas_vendedor(vendedor_id):
    conexion = conectar_db()
    cursor = conexion.cursor()
    ventas = _consultar_ventas(cursor, vendedor_id)
    conexion.close()
    return jsonify(ventas)


@app.route("/api/vendedores/<int:vendedor_id>/ventas", methods=["POST"])
def registrar_venta_vendedor(vendedor_id):
    datos = request.json or {}
    items = datos.get("items") or []
    forma_pago = limpiar_texto(datos.get("forma_pago"))
    observaciones = limpiar_texto(datos.get("observaciones"))

    cliente_id = datos.get("cliente_id")
    if cliente_id in ("", None):
        cliente_id = None
    else:
        try:
            cliente_id = int(cliente_id)
        except (TypeError, ValueError):
            return jsonify({"mensaje": "Cliente inválido"}), 400

    if not forma_pago:
        return jsonify({"mensaje": "La forma de pago es obligatoria"}), 400

    if not isinstance(items, list) or len(items) == 0:
        return jsonify({"mensaje": "La venta debe tener al menos un producto"}), 400

    es_cuenta_corriente = forma_pago.strip().lower() == "cuenta corriente"

    if cliente_id is None:
        return jsonify({
            "mensaje": "Seleccioná un cliente para registrar la venta"
        }), 400

    conexion = conectar_db()
    cursor = conexion.cursor()

    try:
        cursor.execute("""
            SELECT id, nombre
            FROM vendedores
            WHERE id = ? AND activo = 1
        """, (vendedor_id,))
        vendedor = cursor.fetchone()

        if not vendedor:
            conexion.close()
            return jsonify({"mensaje": "Vendedor no encontrado o inactivo"}), 404

        cliente = None
        if cliente_id is not None:
            cursor.execute("""
                SELECT id, nombre, saldo_actual
                FROM clientes
                WHERE id = ?
            """, (cliente_id,))
            cliente = cursor.fetchone()

            if not cliente:
                conexion.close()
                return jsonify({"mensaje": "Cliente no encontrado"}), 404

        detalles_preparados = []
        total = 0.0

        for item in items:
            try:
                producto_id = int(item.get("producto_id"))
                cantidad = int(item.get("cantidad"))
            except (TypeError, ValueError, AttributeError):
                conexion.close()
                return jsonify({"mensaje": "Hay un producto inválido en la venta"}), 400

            if cantidad <= 0:
                conexion.close()
                return jsonify({"mensaje": "Las cantidades deben ser mayores a cero"}), 400

            cursor.execute("""
                SELECT
                    sv.cantidad AS stock_viaje,
                    p.id,
                    p.codigo,
                    p.nombre,
                    p.precio_venta
                FROM stock_viaje sv
                INNER JOIN productos p ON p.id = sv.producto_id
                WHERE sv.vendedor_id = ?
                  AND sv.producto_id = ?
            """, (vendedor_id, producto_id))
            producto = cursor.fetchone()

            if not producto:
                conexion.close()
                return jsonify({
                    "mensaje": "Uno de los productos no pertenece al stock en viaje"
                }), 400

            stock_disponible = int(producto["stock_viaje"] or 0)

            if stock_disponible < cantidad:
                conexion.close()
                return jsonify({
                    "mensaje": (
                        f"Stock insuficiente de {producto['nombre']}. "
                        f"Disponible: {stock_disponible}"
                    )
                }), 400

            precio_unitario = round(float(producto["precio_venta"] or 0), 2)
            subtotal = round(precio_unitario * cantidad, 2)
            total = round(total + subtotal, 2)

            detalles_preparados.append({
                "producto_id": producto_id,
                "cantidad": cantidad,
                "precio_unitario": precio_unitario,
                "subtotal": subtotal,
                "stock_anterior": stock_disponible,
                "nombre": producto["nombre"]
            })

        fecha = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        cursor.execute("""
            INSERT INTO ventas (
                vendedor_id, cliente_id, fecha,
                forma_pago, total, observaciones
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            vendedor_id,
            cliente_id,
            fecha,
            forma_pago,
            total,
            observaciones
        ))
        venta_id = cursor.lastrowid

        for detalle in detalles_preparados:
            nuevo_stock_viaje = detalle["stock_anterior"] - detalle["cantidad"]

            cursor.execute("""
                INSERT INTO venta_detalles (
                    venta_id, producto_id, cantidad,
                    precio_unitario, subtotal
                )
                VALUES (?, ?, ?, ?, ?)
            """, (
                venta_id,
                detalle["producto_id"],
                detalle["cantidad"],
                detalle["precio_unitario"],
                detalle["subtotal"]
            ))

            cursor.execute("""
                UPDATE stock_viaje
                SET cantidad = ?, actualizado_en = ?
                WHERE vendedor_id = ? AND producto_id = ?
            """, (
                nuevo_stock_viaje,
                fecha,
                vendedor_id,
                detalle["producto_id"]
            ))

            cursor.execute("""
                SELECT stock FROM productos WHERE id = ?
            """, (detalle["producto_id"],))
            stock_central = int(cursor.fetchone()["stock"] or 0)

            cursor.execute("""
                INSERT INTO movimientos_stock_viaje (
                    vendedor_id, producto_id, fecha, tipo, cantidad,
                    descripcion, stock_viaje_resultante, stock_central_resultante
                )
                VALUES (?, ?, ?, 'Venta', ?, ?, ?, ?)
            """, (
                vendedor_id,
                detalle["producto_id"],
                fecha,
                -detalle["cantidad"],
                f"Venta #{venta_id}",
                nuevo_stock_viaje,
                stock_central
            ))

        if es_cuenta_corriente and cliente is not None:
            saldo_anterior = float(cliente["saldo_actual"] or 0)
            nuevo_saldo = round(saldo_anterior + total, 2)

            cursor.execute("""
                INSERT INTO cuentas_corrientes (
                    cliente_id, fecha, tipo, descripcion, comprobante,
                    debe, haber, saldo, medio_pago
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                cliente_id,
                fecha,
                "Venta vendedor",
                f"Venta #{venta_id} - {vendedor['nombre']}",
                str(venta_id),
                total,
                0,
                nuevo_saldo,
                "Cuenta corriente"
            ))

            cursor.execute("""
                UPDATE clientes
                SET saldo_actual = ?
                WHERE id = ?
            """, (nuevo_saldo, cliente_id))

        conexion.commit()
    except Exception:
        conexion.rollback()
        conexion.close()
        raise

    conexion.close()

    return jsonify({
        "mensaje": "Venta registrada correctamente",
        "venta_id": venta_id,
        "total": total
    }), 201


@app.route("/api/importar-catalogo", methods=["POST"])
def importar_catalogo():
    if "archivo" not in request.files:
        return jsonify({
            "mensaje": "No se recibió ningún archivo."
        }), 400

    archivo = request.files["archivo"]
    tipo_importacion = request.form.get("tipo_importacion", "catalogo_completo")

    tipos_permitidos = [
        "catalogo_venta",
        "lista_proveedores",
        "catalogo_completo"
    ]

    if tipo_importacion not in tipos_permitidos:
        return jsonify({
            "mensaje": "Tipo de importación no válido."
        }), 400

    if archivo.filename == "":
        return jsonify({
            "mensaje": "Seleccioná un archivo."
        }), 400

    try:
        filas = leer_archivo_importacion(archivo)

        conexion = conectar_db()
        cursor = conexion.cursor()

        proveedores_creados = 0
        productos_creados = 0
        productos_actualizados = 0
        costos_actualizados = 0
        errores = []

        for indice, fila in enumerate(filas, start=2):
            proveedor_nombre = limpiar_texto(
                obtener_valor(fila, "Proveedor")
            )

            codigo = limpiar_texto(
                obtener_valor(fila, "Código", "Codigo", "Cod", "Cod.")
            )

            nombre_producto = limpiar_texto(
                obtener_valor(fila, "Producto", "Nombre", "Artículo", "Articulo", "Descripcion", "Descripción")
            )

            if not nombre_producto:
                errores.append(f"Fila {indice}: falta el nombre del producto.")
                continue

            if not proveedor_nombre:
                if tipo_importacion == "catalogo_venta":
                    proveedor_nombre = "ALMADA / CATÁLOGO DE VENTA"
                else:
                    proveedor_nombre = "SIN PROVEEDOR"

            cursor.execute("""
                SELECT id FROM proveedores
                WHERE LOWER(nombre) = LOWER(?)
            """, (proveedor_nombre,))

            proveedor = cursor.fetchone()

            if proveedor:
                proveedor_id = proveedor["id"]
            else:
                cursor.execute("""
                    INSERT INTO proveedores (
                        nombre, telefono, direccion, localidad, contacto, observaciones
                    )
                    VALUES (?, ?, ?, ?, ?, ?)
                """, (
                    proveedor_nombre,
                    "",
                    "",
                    "",
                    "",
                    "Creado por importación"
                ))

                proveedor_id = cursor.lastrowid
                proveedores_creados += 1

            descripcion = limpiar_texto(
                obtener_valor(fila, "Descripción", "Descripcion", "Detalle")
            )

            if not descripcion:
                descripcion = nombre_producto

            categoria = limpiar_texto(
                obtener_valor(fila, "Categoría", "Categoria", "Rubro", "Linea", "Línea")
            )

            cantidad_caja = int(limpiar_numero(
                obtener_valor(fila, "Cantidad por caja", "Cantidad_caja", "Caja", "Unidades por caja"),
                1
            ))

            if cantidad_caja <= 0:
                cantidad_caja = 1

            precio_venta_excel = limpiar_numero(
                obtener_valor(
                    fila,
                    "Precio venta",
                    "Precio_venta",
                    "Precio unitario",
                    "Precio unidad",
                    "Precio",
                    "Venta"
                ),
                0
            )

            precio_caja_excel = limpiar_numero(
                obtener_valor(fila, "Precio caja", "Precio_caja"),
                0
            )

            precio_base = limpiar_numero(
                obtener_valor(
                    fila,
                    "Precio costo",
                    "Precio_costo",
                    "Precio base",
                    "Precio_base",
                    "Costo",
                    "Compra"
                ),
                0
            )

            bonif1 = limpiar_numero(
                obtener_valor(fila, "Bonificación 1", "Bonificacion 1", "Bonif 1", "Bonif1"),
                0
            )

            bonif2 = limpiar_numero(
                obtener_valor(fila, "Bonificación 2", "Bonificacion 2", "Bonif 2", "Bonif2"),
                0
            )

            bonif3 = limpiar_numero(
                obtener_valor(fila, "Bonificación 3", "Bonificacion 3", "Bonif 3", "Bonif3"),
                0
            )

            iva = limpiar_numero(
                obtener_valor(fila, "IVA", "Iva"),
                21
            )

            flete = limpiar_numero(
                obtener_valor(fila, "Flete"),
                0
            )

            ganancia = limpiar_numero(
                obtener_valor(fila, "Ganancia", "Margen"),
                40
            )

            stock = int(limpiar_numero(
                obtener_valor(fila, "Stock", "Existencia", "Cantidad"),
                0
            ))

            if precio_base > 0:
                precio_final_compra, precio_venta_calculado = calcular_precios(
                    precio_base,
                    bonif1,
                    bonif2,
                    bonif3,
                    iva,
                    flete,
                    ganancia
                )
            else:
                precio_final_compra = 0
                precio_venta_calculado = 0

            precio_venta_final = precio_venta_excel

            if precio_venta_final <= 0 and precio_venta_calculado > 0:
                precio_venta_final = precio_venta_calculado

            precio_unidad = precio_venta_final

            if precio_caja_excel > 0:
                precio_caja = precio_caja_excel
            else:
                precio_caja = precio_unidad * cantidad_caja

            cursor.execute("""
                SELECT id FROM productos
                WHERE codigo = ?
                AND codigo != ''
            """, (codigo,))

            producto_existente = cursor.fetchone()

            if not producto_existente:
                cursor.execute("""
                    SELECT id FROM productos
                    WHERE LOWER(nombre) = LOWER(?)
                    AND proveedor_id = ?
                """, (
                    nombre_producto,
                    proveedor_id
                ))

                producto_existente = cursor.fetchone()

            if producto_existente:
                producto_id = producto_existente["id"]

                cursor.execute("""
                    UPDATE productos
                    SET proveedor_id = ?,
                        nombre = ?,
                        descripcion = ?,
                        categoria = ?,
                        cantidad_caja = ?,
                        precio_unidad = ?,
                        precio_caja = ?,
                        precio_venta = ?,
                        stock = ?
                    WHERE id = ?
                """, (
                    proveedor_id,
                    nombre_producto,
                    descripcion,
                    categoria,
                    cantidad_caja,
                    precio_unidad,
                    precio_caja,
                    precio_venta_final,
                    stock,
                    producto_id
                ))

                productos_actualizados += 1
            else:
                cursor.execute("""
                    INSERT INTO productos (
                        proveedor_id, codigo, nombre, descripcion, categoria, foto,
                        cantidad_caja, precio_unidad, precio_caja,
                        precio_venta, stock
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    proveedor_id,
                    codigo,
                    nombre_producto,
                    descripcion,
                    categoria,
                    "",
                    cantidad_caja,
                    precio_unidad,
                    precio_caja,
                    precio_venta_final,
                    stock
                ))

                producto_id = cursor.lastrowid
                productos_creados += 1

            if tipo_importacion in ["lista_proveedores", "catalogo_completo"]:
                cursor.execute("""
                    SELECT id FROM precios_proveedor
                    WHERE proveedor_id = ?
                    AND producto_id = ?
                """, (
                    proveedor_id,
                    producto_id
                ))

                precio_existente = cursor.fetchone()

                if precio_existente:
                    cursor.execute("""
                        UPDATE precios_proveedor
                        SET precio_base = ?,
                            bonificacion_1 = ?,
                            bonificacion_2 = ?,
                            bonificacion_3 = ?,
                            iva = ?,
                            flete = ?,
                            precio_final_compra = ?,
                            ganancia = ?,
                            precio_venta = ?,
                            fecha_actualizacion = ?
                        WHERE id = ?
                    """, (
                        precio_base,
                        bonif1,
                        bonif2,
                        bonif3,
                        iva,
                        flete,
                        precio_final_compra,
                        ganancia,
                        precio_venta_final,
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        precio_existente["id"]
                    ))
                else:
                    cursor.execute("""
                        INSERT INTO precios_proveedor (
                            proveedor_id, producto_id, precio_base,
                            bonificacion_1, bonificacion_2, bonificacion_3,
                            iva, flete, precio_final_compra, ganancia,
                            precio_venta, fecha_actualizacion
                        )
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        proveedor_id,
                        producto_id,
                        precio_base,
                        bonif1,
                        bonif2,
                        bonif3,
                        iva,
                        flete,
                        precio_final_compra,
                        ganancia,
                        precio_venta_final,
                        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    ))

                costos_actualizados += 1

        conexion.commit()
        conexion.close()

        nombre_tipo = {
            "catalogo_venta": "Catálogo de venta",
            "lista_proveedores": "Lista proveedores / costos",
            "catalogo_completo": "Catálogo completo"
        }

        return jsonify({
            "mensaje": "Importación finalizada.",
            "tipo_importacion": nombre_tipo.get(tipo_importacion, tipo_importacion),
            "proveedores_creados": proveedores_creados,
            "productos_creados": productos_creados,
            "productos_actualizados": productos_actualizados,
            "costos_actualizados": costos_actualizados,
            "errores": errores
        })

    except Exception as error:
        print("ERROR IMPORTAR CATALOGO:", error)

        return jsonify({
            "mensaje": "No se pudo importar el catálogo.",
            "error": str(error)
        }), 500
if __name__ == "__main__":
    crear_tablas()
    asegurar_columnas()
    cargar_datos_iniciales()
    app.run(
        debug=os.environ.get("FLASK_DEBUG", "0") == "1",
        host=os.environ.get("ALMADA2_HOST", "0.0.0.0"),
        port=int(os.environ.get("PORT", "5000"))
    )
